#!/usr/bin/env python3
"""Export normalized slide payloads from a PPTX file."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any

def _normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip().lower()


def _hash_text(text: str) -> str:
    return hashlib.sha1(_normalize_text(text).encode("utf-8")).hexdigest()


def _extract_notes(pptx_path: Path) -> dict[int, str]:
    notes_by_slide: dict[int, str] = {}
    try:
        from pptx import Presentation
    except ImportError:
        return notes_by_slide

    prs = Presentation(str(pptx_path))
    for index, slide in enumerate(prs.slides, start=1):
        notes_text = ""
        try:
            notes_slide = slide.notes_slide
            if notes_slide and notes_slide.notes_text_frame:
                notes_text = (notes_slide.notes_text_frame.text or "").strip()
        except Exception:
            notes_text = ""
        notes_by_slide[index] = notes_text
    return notes_by_slide


def export_pptx_slide_payload(pptx_path: Path, strategy: str = "text") -> dict[str, Any]:
    pptx_path = Path(pptx_path).resolve()
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("python-pptx is required for PPTX export") from exc

    prs = Presentation(str(pptx_path))
    notes_by_slide = _extract_notes(pptx_path)

    slides: list[dict[str, Any]] = []
    for slide_number, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = (shape.text or "").strip()
                if text:
                    texts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        texts.append(row_text)

        notes = (notes_by_slide.get(slide_number) or "").strip()
        content = "\n".join(texts).strip()
        title = texts[0] if texts else f"Slide {slide_number}"
        combined = "\n".join(part for part in [content, notes] if part).strip()
        slides.append(
            {
                "slide_number": slide_number,
                "title": title,
                "content": content,
                "notes": notes,
                "word_count": len(combined.split()) if combined else 0,
                "char_count": len(combined),
                "content_hash": _hash_text(combined),
                "normalized_text": _normalize_text(combined),
            }
        )

    return {
        "source_path": str(pptx_path),
        "filename": pptx_path.name,
        "strategy": strategy,
        "slide_count": len(slides),
        "slides": slides,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Export normalized PPTX slides as JSON.")
    parser.add_argument("pptx_path", help="Path to the PPTX file to export")
    parser.add_argument("--strategy", default="text", choices=["text", "hybrid", "pdf"])
    parser.add_argument("--output", help="Optional output path for JSON payload")
    args = parser.parse_args()

    payload = export_pptx_slide_payload(Path(args.pptx_path), strategy=args.strategy)

    if args.output:
        output_path = Path(args.output)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    else:
        print(json.dumps(payload, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
