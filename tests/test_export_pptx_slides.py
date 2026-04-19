from pathlib import Path

from pptx import Presentation

from scripts.document_processing.export_pptx_slides import export_pptx_slide_payload


def _build_sample_pptx(path: Path) -> None:
    prs = Presentation()

    title_layout = prs.slide_layouts[0]
    bullet_layout = prs.slide_layouts[1]

    slide1 = prs.slides.add_slide(title_layout)
    slide1.shapes.title.text = "Macro Curriculum Blueprint"
    slide1.placeholders[1].text = "A schoolwide coherence frame"
    slide1.notes_slide.notes_text_frame.text = (
        "Presenter note: connect blueprint language to transfer goals."
    )

    slide2 = prs.slides.add_slide(bullet_layout)
    slide2.shapes.title.text = "Learning Ecosystem"
    slide2.placeholders[1].text = "Agency\nCoherence\nSchedule design"
    slide2.notes_slide.notes_text_frame.text = "Speaker note about learner agency."

    prs.save(path)


def test_export_pptx_slide_payload_includes_notes_and_hashes(tmp_path: Path):
    pptx_path = tmp_path / "sample_deck.pptx"
    _build_sample_pptx(pptx_path)

    payload = export_pptx_slide_payload(pptx_path, strategy="text")

    assert payload["slide_count"] == 2
    assert payload["slides"][0]["title"] == "Macro Curriculum Blueprint"
    assert payload["slides"][0]["notes"] == (
        "Presenter note: connect blueprint language to transfer goals."
    )
    assert payload["slides"][1]["word_count"] > 0
    assert payload["slides"][0]["content_hash"]
